"""Export the main KPI tables to a formatted Excel pack (reports/olist_kpi_pack.xlsx)."""
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from kpis import ROOT, findings, load_kpis

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
PERCENT_COLS = ("rate", "ratio", "share")
MONEY_COLS = ("gmv", "brl")
DAY_COLS = ("days", "review")


def number_format(col: str) -> str:
    if any(t in col for t in PERCENT_COLS):
        return "0.0%"
    if any(t in col for t in MONEY_COLS):
        return "#,##0"
    if any(t in col for t in DAY_COLS):
        return "0.00"
    return "#,##0"


def write_sheet(writer, name: str, df: pd.DataFrame):
    df.to_excel(writer, sheet_name=name, index=False)
    ws = writer.sheets[name]
    for i, col in enumerate(df.columns, 1):
        letter = get_column_letter(i)
        cell = ws.cell(row=1, column=i)
        cell.value = col.replace("_", " ").title()
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[letter].width = max(12, min(40, df[col].astype(str).str.len().max() + 2))
        if pd.api.types.is_numeric_dtype(df[col]):
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = number_format(col)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    k = load_kpis()
    out = ROOT / "reports" / "olist_kpi_pack.xlsx"
    sheets = {
        "Monthly trend": k["monthly_trend"],
        "Seller scorecard": k["seller_scorecard"],
        "Lane performance": k["lane_performance"],
        "Customer states": k["customer_state"],
        "Categories": k["category_performance"],
        "Rating vs lateness": k["review_vs_lateness"],
    }
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        summary = pd.DataFrame(findings(k), columns=["finding", "detail"])
        summary.to_excel(writer, sheet_name="Summary", index=False)
        ws = writer.sheets["Summary"]
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 42, 120
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for name, df in sheets.items():
            write_sheet(writer, name, df)
    print(f"wrote {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
